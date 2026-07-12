from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from .config import MAX_EXTRACTED_CHARS, MAX_FILE_BYTES, MAX_PDF_PAGES
from .errors import RuviaError


@dataclass(frozen=True)
class ParsedDocument:
    document_id: str
    filename: str
    text: str
    unit_count: int
    character_count: int
    parser_warnings: list[str]


class DocumentParser:
    def parse(self, filename: str, content_type: str | None, content: bytes) -> ParsedDocument:
        if not content:
            raise RuviaError("empty_file", f"{filename} is empty.")
        if len(content) > MAX_FILE_BYTES:
            raise RuviaError("file_too_large", f"{filename} exceeds the 5 MB file limit.")

        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf":
            return self._parse_pdf(filename, content)
        if suffix == ".docx":
            return self._parse_docx(filename, content)
        raise RuviaError("unsupported_file_type", f"{filename} must be a PDF or DOCX file.")

    def _parse_pdf(self, filename: str, content: bytes) -> ParsedDocument:
        if not content.startswith(b"%PDF"):
            raise RuviaError("disguised_file", f"{filename} does not appear to be a valid PDF.")
        try:
            reader = PdfReader(BytesIO(content))
        except PdfReadError:
            raise RuviaError("unreadable_file", f"{filename} could not be read as a PDF.")
        except Exception:
            raise RuviaError("unreadable_file", f"{filename} could not be read as a PDF.")

        if reader.is_encrypted:
            raise RuviaError("password_protected_pdf", f"{filename} appears to be password protected.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise RuviaError("too_many_pages", f"{filename} exceeds the 20 page limit.")

        chunks: list[str] = []
        warnings: list[str] = []
        for index, page in enumerate(reader.pages):
            try:
                chunks.append(page.extract_text() or "")
            except Exception:
                warnings.append(f"page_{index + 1}_extract_failed")
        text = "\n".join(chunk.strip() for chunk in chunks if chunk.strip()).strip()
        return self._build_document(filename, text, len(reader.pages), warnings)

    def _parse_docx(self, filename: str, content: bytes) -> ParsedDocument:
        if not content.startswith(b"PK"):
            raise RuviaError("disguised_file", f"{filename} does not appear to be a valid DOCX.")
        try:
            with ZipFile(BytesIO(content)) as archive:
                if "[Content_Types].xml" not in archive.namelist():
                    raise RuviaError("disguised_file", f"{filename} does not appear to be a valid DOCX.")
            document = Document(BytesIO(content))
        except RuviaError:
            raise
        except (BadZipFile, ValueError):
            raise RuviaError("unreadable_file", f"{filename} could not be read as a DOCX.")
        except Exception:
            raise RuviaError("unreadable_file", f"{filename} could not be read as a DOCX.")

        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        text = "\n".join(paragraphs).strip()
        return self._build_document(filename, text, len(document.paragraphs), [])

    def _build_document(
        self, filename: str, text: str, unit_count: int, warnings: list[str]
    ) -> ParsedDocument:
        if not text:
            raise RuviaError("empty_extracted_text", f"{filename} did not contain readable text.")
        if len(text) > MAX_EXTRACTED_CHARS:
            raise RuviaError("too_much_text", f"{filename} exceeds the 50,000 character extraction limit.")
        return ParsedDocument(
            document_id=str(uuid4()),
            filename=filename,
            text=text,
            unit_count=unit_count,
            character_count=len(text),
            parser_warnings=warnings,
        )
